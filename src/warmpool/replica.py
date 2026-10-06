from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats as sps

from warmpool.parse import CLIMDIV_STATES, SEASONS

SOUTH = {"AL", "AZ", "FL", "GA", "LA", "MS", "NM", "SC", "TX"}
NORTH = {"ID", "IA", "ME", "MI", "MN", "MT", "ND", "NH", "OR", "SD", "VT", "WA", "WI", "WY"}
NORTHWEST = {"ID", "MT", "OR", "WA"}
ELECTRIC_HEAT = {
    "AL",
    "AR",
    "FL",
    "GA",
    "KY",
    "LA",
    "MS",
    "NC",
    "OK",
    "SC",
    "TN",
    "TX",
    "VA",
    "WA",
    "OR",
    "ID",
}
RETAIL_STATES = sorted(CLIMDIV_STATES + ["DC", "AK", "HI"])
SUPER_EVENTS = {1877: 2.6, 1888: 2.1, 1982: 2.4, 1991: 2.1, 1997: 2.3, 2015: 2.2}
HUB_NAMES = {
    "MIDC": ("Mid Columbia Peak", "Mid C Peak"),
    "PALOVERDE": ("Palo Verde", "Palo Verde Peak"),
    "SP15": ("SP 15", "SP15 EZ Gen DA LMP Peak"),
    "NP15": ("NP15", "NP15 EZ Gen DA LMP Peak"),
    "MASSHUB": ("NEPOOL", "Nepool MH DA LMP Peak"),
    "PJMWEST": ("PJM West", "PJM WH Real Time Peak"),
    "INDIANA": ("Indiana", "Indiana Hub RT Peak"),
    "ERCOTNORTH": ("", "ERCOT North 345KV Peak"),
}
HUB_START = {"INDIANA": "2006-01-04", "NP15": "2009-04-01", "ERCOTNORTH": "2014-04-14"}
HUB_END = {"ERCOTNORTH": "2019-11-07"}
HUB_HEAT_RATE = {
    "MIDC": 8.0,
    "PALOVERDE": 8.5,
    "SP15": 9.5,
    "NP15": 9.5,
    "MASSHUB": 11.0,
    "PJMWEST": 9.0,
    "INDIANA": 8.8,
    "ERCOTNORTH": 8.6,
}
HUB_REGION_STATES = {
    "MIDC": NORTHWEST,
    "PALOVERDE": {"AZ", "NM", "NV"},
    "SP15": {"CA"},
    "NP15": {"CA"},
    "MASSHUB": {"MA", "CT", "RI"},
    "PJMWEST": {"PA", "OH", "VA"},
    "INDIANA": {"IN", "IL"},
    "ERCOTNORTH": {"TX"},
}
WHOLESALE_HEADER = [
    "Price hub",
    "Trade date",
    "Delivery start date",
    "Delivery \nend date",
    "High price $/MWh",
    "Low price $/MWh",
    "Wtd avg price $/MWh",
    "Change",
    "Daily volume MWh",
    "Number of trades",
    "Number of counterparties",
]
NOTE = (
    "The sector, Other, was collected from 1990-2002 and the Transportation sector has been "
    "collected since 2003.\nCustomer counts started to be published in 2007."
)


@dataclass
class Truth:
    heating: dict[str, float] = field(default_factory=dict)
    cooling: dict[str, float] = field(default_factory=dict)
    djf_tmp_per_c: dict[str, float] = field(default_factory=dict)
    billing_weight: float = 0.5
    gas_premium_per_hdd_pct: float = 2.0
    hydro_per_precip_pct: float = -1.0


@dataclass
class World:
    months: pd.DatetimeIndex
    nino: pd.Series
    climate: pd.DataFrame
    truth: Truth


def enso_series(rng: np.random.Generator, end: str) -> pd.Series:
    months = pd.date_range("1869-01-01", end, freq="MS")
    r, theta = 0.965, 2 * np.pi / 50
    phi1, phi2 = 2 * r * np.cos(theta), -(r**2)
    x = np.zeros(months.size)
    noise = rng.normal(0, 0.075, months.size)
    for t in range(2, months.size):
        x[t] = phi1 * x[t - 1] + phi2 * x[t - 2] + noise[t]
    x = x / x.std() * 0.75
    steps = np.arange(months.size)
    for year, peak in SUPER_EVENTS.items():
        center = months.get_loc(pd.Timestamp(year, 12, 1))
        x += (peak - 0.3) * np.exp(-0.5 * ((steps - center) / 4.0) ** 2)
    ramp = (months >= "2026-01-01").astype(float) * np.clip((months.month - 3) / 5.0, -0.2, 1.2)
    x = np.where(months >= "2026-01-01", ramp * 1.45, x)
    return pd.Series(x, index=months)


def _expected_below(threshold: float, mean: np.ndarray, sd: np.ndarray) -> np.ndarray:
    z = (threshold - mean) / sd
    return (threshold - mean) * sps.norm.cdf(z) + sd * sps.norm.pdf(z)


def climate(rng: np.random.Generator, nino: pd.Series, end: str, truth: Truth) -> pd.DataFrame:
    months = pd.date_range("1895-01-01", end, freq="MS")
    djf = nino.rolling(3, center=True).mean()
    rows = []
    for k, state in enumerate(CLIMDIV_STATES):
        lat = 30 + 18 * (k % 12) / 11
        annual = 75 - 0.9 * (lat - 30)
        amp = 12 + 0.6 * (lat - 30)
        tele = 1.6 if state in NORTH else -1.1 if state in SOUTH else 0.2
        truth.djf_tmp_per_c[state] = tele
        wet = -0.12 if state in NORTHWEST else 0.15 if state in SOUTH else 0.0
        seasonal = annual - amp * np.cos(2 * np.pi * (months.month - 1) / 12)
        trend = 0.015 * (months.year - 1895)
        winter = months.month.isin([12, 1, 2])
        signal = np.where(winter, tele * djf.reindex(months).fillna(0).to_numpy(), 0.0)
        tmp = seasonal + trend + signal + rng.normal(0, 2.4, months.size)
        days = months.days_in_month.to_numpy()
        spread = np.where(winter, 8.0, 5.5)
        hdd = days * _expected_below(65.0, tmp, spread)
        cdd = days * _expected_below(-65.0, -tmp, spread)
        base = 3.0 + 1.5 * np.sin(2 * np.pi * (months.month - 3) / 12)
        mult = np.exp(np.where(winter, wet * djf.reindex(months).fillna(0).to_numpy(), 0.0))
        pcpn = rng.gamma(4.0, base * mult / 4.0)
        rows.append(
            pd.DataFrame(
                {"state": state, "date": months, "tmp": tmp, "hdd": hdd, "cdd": cdd, "pcpn": pcpn}
            )
        )
    frame = pd.concat(rows, ignore_index=True)
    frame["year"] = frame["date"].dt.year
    frame["month"] = frame["date"].dt.month
    return frame


def build_world(
    seed: int = 2026, end_climate: str = "2026-08-01", end_enso: str = "2026-09-01"
) -> World:
    rng = np.random.default_rng(seed)
    truth = Truth()
    nino = enso_series(rng, end_enso)
    clim = climate(rng, nino, end_climate, truth)
    for state in RETAIL_STATES:
        truth.heating[state] = 0.013 if state in ELECTRIC_HEAT else 0.006
        truth.cooling[state] = 0.012 if state in SOUTH else 0.008
    return World(months=nino.index, nino=nino, climate=clim, truth=truth)


def write_enso(world: World, folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    smooth = world.nino.rolling(3, center=True).mean().dropna()
    smooth = smooth[smooth.index >= "1950-01-01"]
    roni = ["SEAS   YR  ANOM"]
    oni = [" SEAS  YR   TOTAL   ANOM"]
    for date, value in smooth.items():
        season = SEASONS[date.month - 1]
        warm = 0.01 * (date.year - 1990)
        total = 27.0 + 1.2 * np.sin(2 * np.pi * (date.month - 4) / 12) + value + warm
        roni.append(f"{season}  {date.year} {value:5.2f}")
        oni.append(f"  {season} {date.year} {total:6.2f} {value + warm:6.2f}")
    (folder / "RONI.ascii.txt").write_text("\n".join(roni) + "\n")
    (folder / "oni.ascii.txt").write_text("\n".join(oni) + "\n")
    long = world.nino.copy()
    long.index = long.index
    trend = 0.006 * (long.index.year - 1995)
    long = long + trend
    long = long[long.index >= "1870-01-01"]
    last = long.index.max()
    lines = [f"          1870        {last.year}"]
    for year in range(1870, last.year + 1):
        vals = []
        for month in range(1, 13):
            key = pd.Timestamp(year, month, 1)
            ok = key in long.index and key <= last - pd.DateOffset(months=1)
            vals.append(f"{long[key]:8.2f}" if ok else "  -99.99")
        lines.append(f" {year} " + " ".join(vals))
    lines += [
        "  -99.99",
        "  NINA34",
        " 5N-5S 170W-120W ",
        " HadISST ",
        "  Anomaly from 1981-2010",
        " https://psl.noaa.gov/data/timeseries/month/",
        "  units=degC",
    ]
    (folder / "nino34.long.anom.data").write_text("\n".join(lines) + "\n")
    weeks = pd.date_range("1981-09-02", world.nino.index.max() + pd.Timedelta(days=50), freq="7D")
    daily = world.nino.resample("D").interpolate()
    rows = [
        " Weekly SST data starts week centered on 2Sept1981",
        "",
        "                Nino1+2      Nino3        Nino34        Nino4",
        " Week          SST SSTA     SST SSTA     SST SSTA     SST SSTA",
    ]
    for week in weeks:
        a = float(daily.get(week, daily.iloc[-1]))
        parts = []
        for base, scale in ((22.0, 1.4), (26.0, 1.1), (27.0, 1.0), (28.5, 0.5)):
            anom = round(scale * a, 1)
            parts.append(f"{base + anom:9.1f}{anom:4.1f}")
        rows.append(f" {week.strftime('%d%b%Y').upper()}" + "".join(parts))
    (folder / "wksst9120.for").write_text("\n".join(rows) + "\n")


def write_climdiv(world: World, folder: Path, stamp: str = "20260904") -> None:
    folder.mkdir(parents=True, exist_ok=True)
    clim = world.climate
    first, last = int(clim["year"].min()), int(clim["year"].max())
    years = np.arange(first, last + 1)
    pos = {st: i for i, st in enumerate(CLIMDIV_STATES)}
    specs = {
        "tmpc": ("02", "tmp", "{:7.2f}", -99.90),
        "pcpn": ("01", "pcpn", "{:7.2f}", -9.99),
        "hddc": ("25", "hdd", "{:6.0f}.", -9999.0),
        "cddc": ("26", "cdd", "{:6.0f}.", -9999.0),
    }
    for element, (code, col, fmt, missing) in specs.items():
        cube = np.full((len(CLIMDIV_STATES), years.size, 12), missing)
        cube[
            clim["state"].map(pos).to_numpy(),
            clim["year"].to_numpy() - first,
            clim["month"].to_numpy() - 1,
        ] = clim[col].to_numpy()
        codes = [(f"{i + 1:03d}", i) for i in range(len(CLIMDIV_STATES))]
        codes += [("050", 1), ("101", 2), ("110", 3), ("201", 4)]
        if element in ("tmpc", "pcpn"):
            codes.insert(len(CLIMDIV_STATES), ("049", 5))
        lines = []
        for num, source in codes:
            for yi, year in enumerate(years):
                vals = "".join(fmt.format(v) for v in cube[source, yi])
                lines.append(f"{num}0{code}{year}{vals}   ")
        (folder / f"climdiv-{element}st-v1.0.0-{stamp}").write_text("\n".join(lines) + "\n")


def retail_frame(world: World, end: str = "2026-07-01", seed: int = 3) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    clim = world.climate.set_index(["state", "date"])
    months = pd.date_range("1990-01-01", end, freq="MS")
    days = months.days_in_month.to_numpy()
    w = world.truth.billing_weight
    rows = []
    for k, state in enumerate(RETAIL_STATES):
        source = "MD" if state == "DC" else state
        if source in CLIMDIV_STATES:
            part = clim.loc[source]
            span = pd.date_range(months[0] - pd.DateOffset(months=1), end, freq="MS")
            hdd = part["hdd"].reindex(span).to_numpy() / span.days_in_month.to_numpy()
            cdd = part["cdd"].reindex(span).to_numpy() / span.days_in_month.to_numpy()
            hb = w * hdd[:-1] + (1 - w) * hdd[1:]
            cb = w * cdd[:-1] + (1 - w) * cdd[1:]
        else:
            hb = 20 + 10 * np.cos(2 * np.pi * (months.month - 1) / 12)
            cb = np.zeros(months.size)
        shocks = rng.normal(0, 0.012, months.year.max() - 1989)
        level = np.log(30000 + 9000 * (k % 7)) + 0.012 * (months.year - 1990)
        level = level + shocks[months.year - 1990]
        season = 0.03 * np.cos(2 * np.pi * (months.month - 1) / 12)
        log_daily = (
            level
            + season
            + world.truth.heating[state] * hb
            + world.truth.cooling[state] * cb
            + rng.normal(0, 0.015, months.size)
        )
        res = np.exp(log_daily) * days
        com = res * 0.9
        ind = res * 0.7
        trn = np.where(months.year >= 2003, res * 0.002, 0.0)
        oth = np.where(months.year <= 2002, res * 0.01, np.nan)
        price = 8.0 + 0.2 * (months.year - 1990) + 0.5 * (k % 5)
        customers = np.where(months.year >= 2007, np.round(res / 0.9), np.nan)
        rows.append(
            pd.DataFrame(
                {
                    "state": state,
                    "date": months,
                    "res": res,
                    "com": com,
                    "ind": ind,
                    "trn": trn,
                    "oth": oth,
                    "price": price,
                    "customers": customers,
                }
            )
        )
    return pd.concat(rows, ignore_index=True)


def _sector_cells(sales, price, customers) -> list:
    if not np.isfinite(sales):
        return [".", ".", ".", "."]
    revenue = round(sales * price / 100.0, 2)
    cust = "." if not np.isfinite(customers) else int(customers)
    return [revenue, round(sales, 2), cust, round(price, 2) if sales > 0 else 0]


def write_retail(world: World, folder: Path) -> None:
    from openpyxl import Workbook

    folder.mkdir(parents=True, exist_ok=True)
    frame = retail_frame(world)
    frame["year"] = frame["date"].dt.year
    frame["month"] = frame["date"].dt.month
    frame = frame.sort_values(["year", "month", "state"], ascending=[False, False, True])
    for name, sheet, years, sectors in (
        (
            "HS861M_1990-2009.xlsx",
            "Monthly",
            (1990, 2009),
            ["RESIDENTIAL", "COMMERCIAL", "INDUSTRIAL", "TRANSPORTATION", "OTHER", "TOTAL"],
        ),
        (
            "sales_revenue.xlsx",
            "Monthly-States",
            (2010, 2100),
            ["RESIDENTIAL", "COMMERCIAL", "INDUSTRIAL", "TRANSPORTATION", "TOTAL"],
        ),
    ):
        book = Workbook()
        ws = book.active
        ws.title = sheet
        top = ["", None, None, None]
        mid = ["", None, None, None]
        head = ["Year", "Month", "State", "Data Status"]
        for sector in sectors:
            top += [sector, None, None, None]
            mid += ["Revenue", "Sales", "Customers", "Price"]
            head += ["Thousand Dollars", "Megawatthours", "Count", "Cents/kWh"]
        for row in (top, mid, head):
            ws.append(row)
        part = frame[frame["year"].between(*years)]
        for r in part.itertuples():
            total = r.res + r.com + r.ind + r.trn + (0 if np.isnan(r.oth) else r.oth)
            values = {
                "RESIDENTIAL": r.res,
                "COMMERCIAL": r.com,
                "INDUSTRIAL": r.ind,
                "TRANSPORTATION": r.trn,
                "OTHER": r.oth,
                "TOTAL": total,
            }
            status = "Preliminary" if r.year >= 2025 else "Final"
            cells = [r.year, r.month, r.state, status]
            for sector in sectors:
                cust = (
                    r.customers
                    if sector in ("RESIDENTIAL", "TOTAL")
                    else (np.nan if np.isnan(r.customers) else round(r.customers / 8))
                )
                cells += _sector_cells(values[sector], r.price, cust)
            ws.append(cells)
        ws.append([NOTE])
        book.create_sheet("US-YTD").append(["Year"])
        book.save(folder / name)


def _gas(world: World, rng: np.random.Generator) -> pd.Series:
    days = pd.bdate_range("1997-01-07", world.nino.index.max() + pd.offsets.MonthEnd(1))
    clim = world.climate
    us = clim.groupby("date")["hdd"].mean()
    normal = us.groupby(us.index.month).transform("mean")
    anom = (100.0 * (us - normal) / normal.clip(lower=1.0)).reindex(days, method="ffill").fillna(0)
    winter = days.month.isin([12, 1, 2]).astype(float)
    walk = np.cumsum(rng.normal(0, 0.012, days.size))
    walk = walk - pd.Series(walk, index=days).rolling(500, min_periods=1).mean().to_numpy() * 0.5
    premium = world.truth.gas_premium_per_hdd_pct / 100.0 * anom.to_numpy() * winter
    return pd.Series(3.5 * np.exp(walk + premium), index=days)


def write_gas(world: World, folder: Path, seed: int = 5) -> pd.Series:
    folder.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(seed)
    gas = _gas(world, rng)
    keep = (gas.index.year >= 2007) | (rng.random(gas.size) < 0.3)
    lines = ["observation_date,DHHNGSP"]
    for day, value, ok in zip(gas.index, gas.to_numpy(), keep, strict=True):
        lines.append(f"{day.date()},{value:.2f}" if ok else f"{day.date()},")
    (folder / "DHHNGSP.csv").write_text("\n".join(lines) + "\n")
    return gas


def wholesale_frame(world: World, gas: pd.Series, seed: int = 7) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    clim = world.climate
    rows = []
    nw = clim[clim["state"].isin(NORTHWEST)].groupby("date")["pcpn"].mean()
    winter_sum = nw[nw.index.month.isin([10, 11, 12, 1, 2, 3])]
    season = winter_sum.index.year - (winter_sum.index.month <= 3)
    totals = winter_sum.groupby(season).sum()
    precip_pct = 100.0 * (totals - totals.mean()) / totals.mean()
    for hub, states in HUB_REGION_STATES.items():
        start = HUB_START.get(hub, "2001-01-02")
        end = HUB_END.get(hub, str(gas.index.max().date()))
        days = pd.bdate_range(start, end)
        hdd = clim[clim["state"].isin(states)].groupby("date")["hdd"].mean()
        normal = hdd.groupby(hdd.index.month).transform("mean")
        anom = (100.0 * (hdd - normal) / normal.clip(lower=1.0)).reindex(days, method="ffill")
        price = HUB_HEAT_RATE[hub] * gas.reindex(days, method="ffill").to_numpy()
        price = price * np.exp(0.01 * anom.fillna(0).to_numpy() * days.month.isin([12, 1, 2]))
        if hub == "MIDC":
            spring = days.month.isin([4, 5, 6, 7])
            lag = precip_pct.reindex(days.year - 1).fillna(0).to_numpy()
            price = price * np.exp(world.truth.hydro_per_precip_pct / 100.0 * lag * spring)
        price = price * np.exp(rng.normal(0, 0.08, days.size))
        keep = rng.random(days.size) < 0.97
        for trade, value, ok in zip(days, price, keep, strict=True):
            if ok:
                rows.append((hub, trade, trade + pd.offsets.BDay(1), round(float(value), 2)))
    out = pd.DataFrame(rows, columns=["hub", "trade", "delivery", "price"])
    out["high"] = (out["price"] * 1.04).round(2)
    out["low"] = (out["price"] * 0.96).round(2)
    return out


def _cells(row, name) -> list:
    return [
        name,
        row.trade.to_pydatetime(),
        row.delivery.to_pydatetime(),
        row.delivery.to_pydatetime(),
        row.high,
        row.low,
        row.price,
        "na",
        4000,
        5,
        4,
    ]


def write_wholesale(world: World, gas: pd.Series, folder: Path) -> None:
    import xlwt
    from openpyxl import Workbook

    folder.mkdir(parents=True, exist_ok=True)
    frame = wholesale_frame(world, gas)
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")

    def xls_bytes(rows: list[list], sheet: str) -> bytes:
        book = xlwt.Workbook()
        ws = book.add_sheet(sheet)
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                if hasattr(value, "year"):
                    ws.write(i, j, value, date_style)
                else:
                    ws.write(i, j, value)
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    header = [h.title() if i else "Price Hub" for i, h in enumerate(WHOLESALE_HEADER)]
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        for hub, (old, _) in HUB_NAMES.items():
            part = frame[(frame["hub"] == hub) & (frame["trade"].dt.year <= 2013)]
            if part.empty or not old:
                continue
            rows = [header] + [_cells(r, old) for r in part.itertuples()]
            zf.writestr(f"{old} Hub.xls", xls_bytes(rows, "Sheet1"))
    (folder / "ice_electric-historical.zip").write_bytes(archive.getvalue())
    last_year = int(frame["trade"].dt.year.max())
    for year in range(2014, last_year + 1):
        part = frame[frame["trade"].dt.year == year].sort_values(["hub", "trade"])
        rows = [WHOLESALE_HEADER] + [_cells(r, HUB_NAMES[r.hub][1]) for r in part.itertuples()]
        if year <= 2016:
            (folder / f"ice_electric-{year}final.xls").write_bytes(xls_bytes(rows, str(year)))
            continue
        book = Workbook()
        ws = book.active
        ws.title = str(year)
        for row in rows:
            ws.append(row)
        name = (
            f"ice_electric-{year}.xlsx" if year == last_year else f"ice_electric-{year}final.xlsx"
        )
        book.save(folder / name)


def write_all(raw: Path, seed: int = 2026) -> World:
    raw = Path(raw)
    world = build_world(seed)
    write_enso(world, raw / "enso")
    write_climdiv(world, raw / "climdiv")
    write_retail(world, raw / "retail")
    gas = write_gas(world, raw / "gas")
    write_wholesale(world, gas, raw / "wholesale")
    return world
